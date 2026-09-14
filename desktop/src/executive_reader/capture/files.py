"""Read documents straight from disk.

This is why the app never has to scroll a PDF. Parsing the file yields every
page at once, in order, with no capture and no recognition error. Screen OCR
is only for pages that have no text layer at all.
"""
from __future__ import annotations

import posixpath
import re
import zipfile
from pathlib import Path
from urllib.parse import unquote

from ..document import Document

PDF_EXT = {".pdf"}
EPUB_EXT = {".epub"}
DOCX_EXT = {".docx"}
TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".json",
            ".py", ".js", ".ts", ".html", ".htm", ".xml", ".yaml", ".yml"}
SUPPORTED = PDF_EXT | EPUB_EXT | DOCX_EXT | TEXT_EXT

_PAGE_NUMBER = re.compile(r"^\s*(?:page\s*)?\d{1,4}\s*(?:/\s*\d{1,4})?\s*$", re.I)


class UnsupportedFile(RuntimeError):
    pass


def is_supported(path) -> bool:
    return Path(path).suffix.lower() in SUPPORTED


def _strip_repeated_lines(pages: list[str], min_pages: int = 3) -> list[str]:
    """Drop running headers and footers.

    A short line that appears on most pages is furniture, not content, and
    hearing it between every page is the fastest way to make a long PDF
    unlistenable.
    """
    # Two jobs, and they used to share one exit. Boilerplate is found
    # statistically and needs several pages to be sure. A page number is found
    # by its shape and needs nothing. Returning early when no boilerplate was
    # detected skipped the page numbers as well, so the same document kept or
    # dropped them depending only on whether it happened to have a running
    # footer. A PDF with numbered pages and no header is the ordinary case, and
    # it read "Page 1", "Page 2" aloud between every page.
    boilerplate: set[str] = set()
    if len(pages) >= min_pages:
        counts: dict[str, int] = {}
        for page in pages:
            lines = [ln.strip() for ln in page.splitlines() if ln.strip()]
            for line in set(lines[:2] + lines[-2:]):
                if len(line) <= 90:
                    counts[line] = counts.get(line, 0) + 1
        threshold = max(2, int(len(pages) * 0.6))
        boilerplate = {line for line, n in counts.items() if n >= threshold}

    cleaned = []
    for page in pages:
        keep = [ln for ln in page.splitlines()
                if ln.strip() not in boilerplate and not _PAGE_NUMBER.match(ln)]
        cleaned.append("\n".join(keep))
    return cleaned


#: Below this much text in total, a document has no usable text layer at all.
#: Deliberately near zero: the question is whether extraction produced
#: anything, not whether it produced a lot. An earlier floor of 60 characters
#: called a one-page note scanned and refused to read text it had extracted
#: perfectly.
MIN_TEXT_CHARS = 16
#: And per page, so a hundred-page scan with one stray caption is still a scan.
MIN_CHARS_PER_PAGE = 8


def looks_scanned(text: str, page_count: int) -> bool:
    """Whether a PDF has no text layer, as opposed to little text.

    A separate function because it is a judgment rather than a detail, and
    because a predicate can be tested at page counts a fixture cannot easily
    reach.
    """
    return len(text.strip()) < max(MIN_TEXT_CHARS, page_count * MIN_CHARS_PER_PAGE)


def read_pdf(path) -> Document:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise UnsupportedFile(
            "pypdfium2 is not installed. Run: pip install pypdfium2") from exc

    path = Path(path)
    pdf = pdfium.PdfDocument(str(path))
    try:
        pages = []
        for i in range(len(pdf)):
            page = pdf[i]
            try:
                textpage = page.get_textpage()
                pages.append(textpage.get_text_bounded() or "")
            except Exception:
                pages.append("")
        page_count = len(pdf)
    finally:
        pdf.close()

    pages = _strip_repeated_lines(pages)
    text = "\n\n".join(p.strip() for p in pages if p.strip())
    scanned = looks_scanned(text, page_count)
    return Document(
        text=text, title=path.stem, uri=path.as_uri(), source="file",
        meta={"path": str(path), "pages": page_count, "kind": "pdf",
              "needs_ocr": scanned},
    )


def _epub_spine(archive: zipfile.ZipFile) -> tuple[list[str], str]:
    """Reading order and title from the package document.

    Chapters must be read in spine order, not archive order: a zip lists its
    entries however they were written, so trusting that order reads a book with
    the chapters shuffled.
    """
    from bs4 import BeautifulSoup

    try:
        container = archive.read("META-INF/container.xml")
    except KeyError as exc:
        raise UnsupportedFile("Not a valid EPUB: no container.xml") from exc

    root = BeautifulSoup(container, "xml").find("rootfile")
    opf_path = root.get("full-path") if root else None
    if not opf_path:
        raise UnsupportedFile("Not a valid EPUB: no package document")

    package = BeautifulSoup(archive.read(opf_path), "xml")
    base = posixpath.dirname(opf_path)

    hrefs: dict[str, str] = {}
    for item in package.find_all("item"):
        item_id, href = item.get("id"), item.get("href")
        if item_id and href:
            # hrefs are relative to the package document, not the archive root.
            hrefs[item_id] = posixpath.normpath(
                posixpath.join(base, unquote(href))) if base else unquote(href)

    order = []
    for ref in package.find_all("itemref"):
        target = hrefs.get(ref.get("idref") or "")
        if target:
            order.append(target)

    title_tag = package.find("title")
    title = title_tag.get_text(strip=True) if title_tag else ""
    return order, title


def read_epub(path) -> Document:
    """Read an EPUB with the standard library and an HTML parser.

    Deliberately not EbookLib: that package is AGPL, and distributing it would
    put copyleft over this whole application. An EPUB is a zip of XHTML with a
    manifest, so the dependency buys convenience rather than capability. See
    CAVEATS.md.
    """
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise UnsupportedFile(
            "beautifulsoup4 is required for EPUB files.") from exc

    path = Path(path)
    chapters: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            order, title = _epub_spine(archive)
            names = set(archive.namelist())
            for href in order:
                if href not in names:
                    continue
                soup = BeautifulSoup(archive.read(href), "html.parser")
                for tag in soup(["script", "style", "nav"]):
                    tag.decompose()
                chunk = soup.get_text("\n").strip()
                if chunk:
                    chapters.append(chunk)
    except zipfile.BadZipFile as exc:
        raise UnsupportedFile("That EPUB is not a readable archive.") from exc

    return Document(text="\n\n".join(chapters), title=title or path.stem,
                    uri=path.as_uri(), source="file",
                    meta={"path": str(path), "kind": "epub",
                          "chapters": len(chapters)})


def read_docx(path) -> Document:
    try:
        import docx
    except ImportError as exc:
        raise UnsupportedFile(
            "python-docx is not installed. Run: pip install python-docx") from exc
    path = Path(path)
    document = docx.Document(str(path))
    parts = [text for text in _docx_blocks(document) if text]
    return Document(text="\n\n".join(parts), title=path.stem, uri=path.as_uri(),
                    source="file", meta={"path": str(path), "kind": "docx"})


def _docx_blocks(document):
    """Paragraphs and tables in the order they appear on the page.

    `document.paragraphs` and `document.tables` are two separate lists, so
    reading one then the other moves every table to the end of the document.
    In a report with a table on page two, the listener hears it after page
    fifty, with nothing to say it moved. The same mistake in the other
    direction is why EPUB reading follows the spine rather than the archive.

    The body's own XML children are the only thing that knows the real order.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document).text.strip()
        elif child.tag == qn("w:tbl"):
            for row in _docx_rows(Table(child, document)):
                yield row


def _docx_rows(table):
    """One line per row, with each merged cell spoken once.

    A cell merged across three columns is returned three times by `row.cells`,
    once per column it covers, so a totals row came out as "Total for the year,
    Total for the year, Total for the year". A vertical merge repeats the same
    way down the rows.

    Every one of those is the same cell and shares one underlying element, so
    identity separates a merge from a table that genuinely repeats a value in
    neighbouring columns. Comparing the text would silently collapse that too.
    """
    seen: set[int] = set()
    # `alive` exists only to hold references, and removing it loses table data
    # silently. lxml builds an element proxy on demand and drops it as soon as
    # nothing points at it, and CPython then hands the same address to the next
    # one, so an id() taken from a transient proxy collides with an unrelated
    # cell a row later. Measured on a three-row table: the body cell "North"
    # was handed the address of the header cell "Q1" and was skipped as a
    # duplicate. The failure deletes content rather than raising, and it moves
    # around with garbage collection, so it is close to unreproducible once it
    # ships.
    alive: list = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            element = getattr(cell, "_tc", None)
            if element is not None:
                if id(element) in seen:
                    continue
                seen.add(id(element))
                alive.append(element)
            text = cell.text.strip()
            if text:
                cells.append(text)
        if cells:
            yield ", ".join(cells)


def read_text(path) -> Document:
    path = Path(path)
    for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            body = path.read_text(encoding=encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        body = path.read_bytes().decode("utf-8", "replace")
    return Document(text=body, title=path.stem, uri=path.as_uri(), source="file",
                    meta={"path": str(path), "kind": path.suffix.lstrip(".")})


def read(path) -> Document:
    """Parse any supported file. Raises UnsupportedFile for anything else."""
    path = Path(path)
    if not path.exists():
        raise UnsupportedFile("File not found: " + str(path))
    ext = path.suffix.lower()
    if ext in PDF_EXT:
        return read_pdf(path)
    if ext in EPUB_EXT:
        return read_epub(path)
    if ext in DOCX_EXT:
        return read_docx(path)
    if ext in TEXT_EXT:
        return read_text(path)
    raise UnsupportedFile("Cannot read " + ext + " files")
